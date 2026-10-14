"""
export.py
- PDF:  the rendered handwriting pages, at true A4/A5/Letter size.
- DOCX: the same document as *editable* Word text (real paragraphs, headings,
        pictures, spacing and margins), set in the chosen font family. It only
        looks handwritten if that font is installed where the file is opened.
"""
import io
from typing import Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

from .handwriting import DPI, process_image
from .markup import parse_markup
from .schemas import RenderSettings

PAGE_MM = {"a4": (210, 297), "a5": (148, 210), "letter": (215.9, 279.4)}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def pages_to_pdf(pages: List[Image.Image], out_path: str) -> str:
    if not pages:
        raise ValueError("No pages to export")
    rgb = [p.convert("RGB") for p in pages]
    # resolution tells the PDF the physical page size (otherwise it is huge)
    rgb[0].save(out_path, save_all=True, append_images=rgb[1:], resolution=float(DPI))
    return out_path


def markup_to_docx(text: str, images: Dict[str, Image.Image], s: RenderSettings,
                   out_path: str, font_family: str) -> str:
    doc = Document()
    background = OxmlElement("w:background")
    background.set(qn("w:color"), s.paper_color.lstrip("#").upper())
    doc._element.insert(0, background)
    doc.settings.element.append(OxmlElement("w:displayBackgroundShape"))
    sec = doc.sections[0]
    w_mm, h_mm = PAGE_MM[s.page_size]
    sec.page_width, sec.page_height = Mm(w_mm), Mm(h_mm)
    sec.top_margin, sec.bottom_margin = Mm(s.margin_top_mm), Mm(s.margin_bottom_mm)
    sec.left_margin, sec.right_margin = Mm(s.margin_left_mm), Mm(s.margin_right_mm)
    text_width_mm = w_mm - s.margin_left_mm - s.margin_right_mm

    pt = round(s.font_size * 72 / DPI, 1)
    ink = RGBColor(*(int(s.ink_color[i:i + 2], 16) for i in (1, 3, 5)))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font_family, Pt(pt)

    def style_run(run, size=pt, bold=False):
        run.font.name, run.font.size, run.bold = font_family, Pt(size), bold
        run.font.color.rgb = ink

    def fmt(par, gap=True):
        f = par.paragraph_format
        f.line_spacing = max(1.0, s.line_spacing / 1.2)  # Word's "single" is ~1.2 em
        f.space_after = Pt(pt * s.paragraph_gap) if gap else Pt(0)
        par.alignment = ALIGN[s.align]

    for b in parse_markup(text):
        if b.kind == "paragraph":
            par = doc.add_paragraph()
            fmt(par)
            par.paragraph_format.first_line_indent = Mm(s.first_line_indent_mm)
            for i, line in enumerate(b.lines):
                run = par.add_run(line)
                style_run(run)
                if i < len(b.lines) - 1:
                    run.add_break()
        elif b.kind == "heading":
            par = doc.add_paragraph()
            fmt(par)
            par.alignment = ALIGN[s.align if s.align in ("center", "right") else "left"]
            par.paragraph_format.keep_with_next = True
            run = par.add_run(b.lines[0])
            style_run(run, size=round(pt * s.heading_scale, 1), bold=True)
            run.underline = True
        elif b.kind == "image":
            raw = images.get(b.image_id)
            img = process_image(raw, s.image_mode, s.ink_color) if raw is not None else None
            if img is None:
                continue
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(buf, width=Mm(text_width_mm * (b.width_pct or s.image_width_pct) / 100))
        elif b.kind == "pagebreak":
            doc.add_page_break()
    doc.save(out_path)
    return out_path
