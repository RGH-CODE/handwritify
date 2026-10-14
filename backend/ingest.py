"""
ingest.py
Reads uploaded files (PDF, PPTX, JPG, PNG) in the order given and turns
them into pages of ordered items (text paragraphs + pictures).

- Photos / scans      -> OCR with layout analysis (ocr.analyze_page):
                         text is read, diagrams/figures are cropped out.
- PDF with text layer -> text blocks are used directly; embedded pictures
                         and vector diagrams are rendered to images.
                         Pages with no text layer are treated as scans.
- PPTX                -> text (titles become headings) and pictures in
                         top-to-bottom order. If a slide has no text at all
                         and LibreOffice is installed, it is rendered and OCR'd.
"""
from __future__ import annotations

import io
import os
import subprocess
import tempfile
from dataclasses import dataclass, field
from typing import List

import pymupdf
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .markup import reflow
from .models import ImageItem, Item, TextItem, flatten_rgb
from .ocr import analyze_page

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".jpg", ".jpeg", ".png"}
MIN_PICTURE_PX = 40


@dataclass
class Page:
    source: str
    page: int
    items: List[Item] = field(default_factory=list)


def ingest_files(filepaths: List[str]) -> List[Page]:
    """filepaths must already be in the order the user wants them combined."""
    pages: List[Page] = []
    for path in filepaths:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            pages.extend(_ingest_pdf(path))
        elif ext == ".pptx":
            pages.extend(_ingest_pptx(path))
        elif ext in (".jpg", ".jpeg", ".png"):
            pages.extend(_ingest_image(path))
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    return pages


# ---------------------------------------------------------------- images
def _ingest_image(path: str) -> List[Page]:
    img = Image.open(path)
    img = ImageOps.exif_transpose(img)  # phone photos carry a rotation flag
    return [Page(os.path.basename(path), 1, analyze_page(flatten_rgb(img)))]


# ------------------------------------------------------------------- PDF
def _overlap_ratio(a: pymupdf.Rect, b: pymupdf.Rect) -> float:
    """Share of rect `a` that lies inside rect `b`."""
    inter = a & b
    if inter.is_empty or a.get_area() == 0:
        return 0.0
    return inter.get_area() / a.get_area()


def _render_clip(page: pymupdf.Page, rect: pymupdf.Rect) -> Image.Image:
    pix = page.get_pixmap(clip=rect, dpi=200)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def _pdf_page_items(pg: pymupdf.Page) -> List[Item]:
    if len(pg.get_text().strip()) < 20:  # scanned page: no usable text layer
        pix = pg.get_pixmap(dpi=200)
        return analyze_page(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))

    page_area = pg.rect.get_area()
    blocks = pg.get_text("blocks", sort=True)

    figures: List[pymupdf.Rect] = []
    for info in pg.get_image_info():  # embedded raster pictures
        r = pymupdf.Rect(info["bbox"]) & pg.rect
        # skip tiny icons and full-page backgrounds (scans with hidden text)
        if r.width >= MIN_PICTURE_PX and r.height >= MIN_PICTURE_PX and r.get_area() < 0.85 * page_area:
            figures.append(r)
    try:  # vector diagrams / tables drawn with lines
        for r in pg.cluster_drawings():
            if r.width < 60 or r.height < 60 or r.get_area() > 0.85 * page_area:
                continue
            if any(_overlap_ratio(r, f) > 0.5 for f in figures):
                continue
            figures.append(r)
    except Exception:
        pass

    entries = []  # (y, x, item)
    for r in figures:
        entries.append((r.y0, r.x0, ImageItem(_render_clip(pg, r))))
    for b in blocks:
        if b[6] != 0:
            continue
        rect = pymupdf.Rect(b[:4])
        if any(_overlap_ratio(rect, f) > 0.6 for f in figures):
            continue  # a label that belongs to a diagram
        text = reflow(b[4])
        if text:
            entries.append((rect.y0, rect.x0, TextItem(text)))
    entries.sort(key=lambda e: (round(e[0]), e[1]))
    return [e[2] for e in entries]


def _ingest_pdf(path: str) -> List[Page]:
    name = os.path.basename(path)
    doc = pymupdf.open(path)
    try:
        return [Page(name, i, _pdf_page_items(pg)) for i, pg in enumerate(doc, start=1)]
    finally:
        doc.close()


# ------------------------------------------------------------------ PPTX
def _flatten_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flatten_shapes(sh.shapes)
        else:
            yield sh


def _is_title(sh) -> bool:
    try:
        return sh.is_placeholder and sh.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    except Exception:
        return False


def _shape_items(sh) -> List[Item]:
    out: List[Item] = []
    try:
        if sh.has_text_frame:
            paras = [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()]
            if paras:
                out.append(TextItem("# " + " ".join(paras) if _is_title(sh) else "\n".join(paras)))
    except Exception:
        pass
    try:
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in sh.table.rows]
            if rows:
                out.append(TextItem("\n".join(rows)))
    except Exception:
        pass
    try:
        blob = sh.image.blob  # Picture / PlaceholderPicture only
        img = flatten_rgb(Image.open(io.BytesIO(blob)))
        if img.width >= MIN_PICTURE_PX and img.height >= MIN_PICTURE_PX:
            out.append(ImageItem(img))
    except Exception:
        pass  # not a picture, or a format PIL can't read (e.g. EMF)
    return out


def _render_pptx_to_images(path: str) -> List[Image.Image]:
    """Best-effort slide rasterisation via LibreOffice; [] if unavailable."""
    try:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180,
            )
            pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
            doc = pymupdf.open(pdf)
            images = []
            for pg in doc:
                pix = pg.get_pixmap(dpi=200)
                images.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            doc.close()
            return images
    except Exception:
        return []


def _ingest_pptx(path: str) -> List[Page]:
    name = os.path.basename(path)
    prs = Presentation(path)
    rendered = None  # only rasterise the deck if some slide needs OCR
    pages: List[Page] = []
    for i, slide in enumerate(prs.slides, start=1):
        shapes = sorted(_flatten_shapes(slide.shapes),
                        key=lambda s: (s.top or 0, s.left or 0))
        items: List[Item] = []
        for sh in shapes:
            items.extend(_shape_items(sh))
        if not any(isinstance(it, TextItem) for it in items):
            if rendered is None:
                rendered = _render_pptx_to_images(path)
            if i - 1 < len(rendered):
                items = analyze_page(rendered[i - 1])
        pages.append(Page(name, i, items))
    return pages
